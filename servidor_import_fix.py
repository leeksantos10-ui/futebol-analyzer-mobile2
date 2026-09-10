import sqlite3, json, os, csv, io
from datetime import datetime, date, time
from pathlib import Path
from http.server import BaseHTTPRequestHandler, HTTPServer
from urllib.parse import urlparse, parse_qs
from email.parser import BytesParser
from email.policy import default

BASE = Path(__file__).resolve().parent
DB = str(BASE / "banco_jogos_15323.sqlite")
PORT = int(os.environ.get("PORT", "8005"))

METRICS = {
    "LG Score Casa":"LG Score Casa", "LG Score Visitante":"LG Score Visitante", "H-Score":"H-Score",
    "Média Gols Feitos Casa":"Media Gols Feitos Casa", "Média Gols Feitos Visitante":"Media Gols Feitos Visitante",
    "CV Gols Feitos Casa":"CV Gols Feitos Casa", "CV Gols Feitos Visitante":"CV Gols Feitos Visitante",
    "Média Gols Sofridos Casa":"Media Gols Sofridos Casa", "Média Gols Sofridos Visitante":"Media Gols Sofridos Visitante",
    "CV Gols Sofridos Casa":"CV Gols Sofridos Casa", "CV Gols Sofridos Visitante":"CV Gols Sofridos Visitante",
    "Média Gols Marcados HT Casa":"Media Gols Marcados HT Casa", "Média Gols Marcados HT Visitante":"Media Gols Marcados HT Visitante",
    "Média Gols Sofridos HT Casa":"Media Gols Sofridos HT Casa", "Média Gols Sofridos HT Visitante":"Media Gols Sofridos HT Visitante",
    "Média Custo do Gol 2,0 Casa":"Media Custo do Gol 2,0 Casa", "Média Custo do Gol 2,0 Visitante":"Media Custo do Gol 2,0 Visitante",
    "CV Custo do Gol 2,0 Casa":"CV Custo do Gol 2,0 Casa", "CV Custo do Gol 2,0 Visitante":"CV Custo do Gol 2,0 Visitante",
    "Jogos Over 1,5 FT Casa":"Jogos Over 1,5 FT Casa", "Jogos Over 1,5 FT Visitante":"Jogos Over 1,5 FT Visitante",
    "Jogos Over 2,5 FT Casa":"Jogos Over 2,5 FT Casa", "Jogos Over 2,5 FT Visitante":"Jogos Over 2,5 FT Visitante",
    "Pontos Casa":"Pontos Casa", "Pontos Visitante":"Pontos Visitante",
    "Odd Casa":"Odd Casa", "Odd Empate":"Odd Empate", "Odd Visitante":"Odd Visitante",
    "Odd Over 0,5 HT":"Over 0,5 HT", "Odd Over 0,5 FT":"Over 0,5 FT", "Odd Over 1,5 FT":"Over 1,5 FT", "Odd Over 2,5 FT":"Over 2,5 FT",
    "Odd BTTS Sim FT":"BTTS Sim FT", "1X ODD":"1x ODD", "2X ODD":"2X ODD",
    "Força Custo do Gol Casa":"FORÇA CUSTO DO GOL CASA", "Força Custo do Gol Visitante":"FORÇA CUSTO DO GOL VISITANTE",
    "Força Ofensiva FT Casa":"FORÇA OFENSIVA FT CASA", "Força Ofensiva FT Visitante":"FORÇA OFENSIVA VISITAN FT",
    "Força Defensiva FT Casa":"FORÇA DEFENSIVA CASA FT", "Força Defensiva FT Visitante":"FORÇA DEFENSIVA VISITAN FT",
    "Força Ofensiva HT Casa":"FORÇA OFENSIVA CASA HT", "Força Ofensiva HT Visitante":"FORÇA OFENSIVA VISITAN HT",
    "Força Defensiva HT Casa":"FORÇA DEFENSIVA CASA HT", "Força Defensiva HT Visitante":"FORÇA DEFENSIVA VISITAN HT",
    "CV Combinado Custo do Gol":"CV COMBINADO CUSTO DO GOL", "CV Combinado Gols Feitos":"CV COMBINADO GOLS FEITO", "CV Combinado Gols Sofridos":"CV COMBINADO GOLS SOFRIDOS"
}

MARKETS = {
    "Over 0,5 FT": ("Over 0,5 FT", lambda h,a: h+a>=1), "Over 1,5 FT": ("Over 1,5 FT", lambda h,a: h+a>=2),
    "Over 2,5 FT": ("Over 2,5 FT", lambda h,a: h+a>=3), "BTTS Sim FT": ("BTTS Sim FT", lambda h,a: h>=1 and a>=1),
    "Casa": ("Odd Casa", lambda h,a: h>a), "Empate": ("Odd Empate", lambda h,a: h==a), "Visitante": ("Odd Visitante", lambda h,a: a>h),
    "1X": ("1x ODD", lambda h,a: h>=a), "2X": ("2X ODD", lambda h,a: a>=h)
}

def fnum(v):
    try:
        if v is None or str(v).strip()=="": return None
        return float(str(v).strip().replace(",","."))
    except: return None

def cmp(v,op,x):
    if v is None or x is None:return False
    return {">":v>x,">=":v>=x,"<":v<x,"<=":v<=x,"=":abs(v-x)<1e-12}.get(op,False)

def stats(rows,stake):
    bal=0; peak=0; dd=0; ddp=0; g=r=0; odds=[]
    for row in rows:
        odd=fnum(row["odd"])
        if odd is None or odd<=0: continue
        if row["green"]: g+=1; bal += stake*(odd-1)
        else: r+=1; bal -= stake
        peak=max(peak,bal); d=peak-bal; dd=max(dd,d)
        if peak>0: ddp=max(ddp,d/peak*100)
        odds.append(odd)
    n=g+r
    return {"entries":n,"greens":g,"reds":r,"winrate":g/n*100 if n else 0,"avgodd":sum(odds)/len(odds) if odds else 0,"lp":bal,"roi":bal/(n*stake)*100 if n else 0,"maxdd":dd,"maxddpct":ddp}

def evaluate(strategy,market,stake):
    if market not in MARKETS: raise ValueError("Mercado inválido")
    conn=sqlite3.connect(DB); conn.row_factory=sqlite3.Row
    cols={"Gols Casa FT","Gols Visitante FT","Data","Hora","Confronto","Campeonato",MARKETS[market][0]}
    for rule in strategy.get("rules",[]):
        if METRICS.get(rule.get("metric")): cols.add(METRICS[rule["metric"]])
        if rule.get("compare_mode")=="metric" and METRICS.get(rule.get("rhs_metric")): cols.add(METRICS[rule["rhs_metric"]])
    sql="SELECT "+",".join('"'+c.replace('"','""')+'"' for c in cols)+" FROM jogos"
    rows=conn.execute(sql).fetchall(); conn.close(); out=[]
    champs=[str(x).strip() for x in strategy.get("championships",[]) if str(x).strip()]
    for row in rows:
        v={c:row[c] for c in cols}
        if champs and str(v.get("Campeonato","")).strip() not in champs: continue
        ok=True
        for rule in strategy.get("rules",[]):
            left=fnum(v.get(METRICS.get(rule.get("metric"))))
            right=fnum(v.get(METRICS.get(rule.get("rhs_metric")))) if rule.get("compare_mode")=="metric" else fnum(rule.get("value"))
            if not cmp(left,rule.get("op"),right): ok=False; break
        if not ok: continue
        h=fnum(v["Gols Casa FT"]); a=fnum(v["Gols Visitante FT"]); odd=fnum(v[MARKETS[market][0]])
        if h is None or a is None or odd is None or odd<=0: continue
        out.append({"data":v["Data"],"hora":v["Hora"],"confronto":v["Confronto"],"campeonato":v["Campeonato"],"odd":odd,"green":bool(MARKETS[market][1](h,a)),"casa":h,"visitante":a})
    valid=[x for x in out if isinstance(x["data"],str) and x["data"] not in ("Data","")]; valid.sort(key=lambda x:(x["data"],str(x["hora"])))
    allst=stats(out,stake); st=stats(valid,stake); bal=0; curve=[]; months={}
    for x in valid:
        bal += stake*(x["odd"]-1) if x["green"] else -stake; curve.append({"data":x["data"],"banca":bal}); m=str(x["data"])[:7]
        z=months.setdefault(m,{"mes":m,"jogos":0,"green":0,"red":0,"lp":0}); z["jogos"]+=1
        if x["green"]: z["green"]+=1; z["lp"]+=stake*(x["odd"]-1)
        else: z["red"]+=1; z["lp"]-=stake
    for z in months.values(): z["winrate"]=z["green"]/z["jogos"]*100 if z["jogos"] else 0; z["roi"]=z["lp"]/(z["jogos"]*stake)*100 if z["jogos"] else 0
    leagues={}
    for x in valid:
        key=str(x.get("campeonato") or "Sem campeonato").strip() or "Sem campeonato"; z=leagues.setdefault(key,{"campeonato":key,"jogos":0,"green":0,"red":0,"lp":0}); z["jogos"]+=1
        if x["green"]: z["green"]+=1; z["lp"]+=stake*(x["odd"]-1)
        else: z["red"]+=1; z["lp"]-=stake
    for z in leagues.values(): z["winrate"]=z["green"]/z["jogos"]*100 if z["jogos"] else 0; z["roi"]=z["lp"]/(z["jogos"]*stake)*100 if z["jogos"] else 0
    st.update({"all":allst,"curve":curve,"monthly":sorted(months.values(),key=lambda x:x["mes"]),"by_league":sorted(leagues.values(),key=lambda x:x["lp"],reverse=True),"matches":[{"data":x["data"],"hora":x["hora"],"confronto":x["confronto"],"campeonato":x["campeonato"],"odd":x["odd"],"resultado":"GREEN" if x["green"] else "RED","placar":f'{int(x["casa"])} x {int(x["visitante"])}'} for x in out]})
    return st

def parse_upload(body, content_type):
    msg = BytesParser(policy=default).parsebytes((f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n").encode()+body)
    for part in msg.iter_attachments():
        filename=part.get_filename() or "arquivo"
        return filename, part.get_payload(decode=True) or b""
    raise ValueError("Arquivo não encontrado no envio")

def normalize_header(x):
    return str(x or "").strip().replace("\ufeff","")

def excel_value(v, header=""):
    # Converte datas/horas do Excel para texto estável, evitando objetos
    # datetime que não podem ser enviados como JSON.
    h = normalize_header(header).lower()
    if isinstance(v, datetime):
        if "hora" in h and "data" not in h:
            return v.strftime("%H:%M")
        if "data" in h:
            return v.strftime("%d/%m/%Y")
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    return v

def json_safe(v):
    # Segurança extra para qualquer valor vindo da planilha.
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat()
    if isinstance(v, dict):
        return {str(k): json_safe(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [json_safe(x) for x in v]
    return v

def read_rows(filename, data):
    ext=filename.lower().rsplit(".",1)[-1]
    if ext=="xlsx":
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        ws=wb.active; it=ws.iter_rows(values_only=True)
        headers=[normalize_header(x) for x in next(it)]
        rows=[]
        for vals in it:
            if not any(v not in (None,"") for v in vals): continue
            row={}
            for i, header in enumerate(headers):
                if not header: continue
                value = vals[i] if i < len(vals) else None
                row[header] = excel_value(value, header)
            rows.append(row)
        return headers, rows
    if ext=="csv":
        text=data.decode("utf-8-sig",errors="replace")
        sample=text[:4096]
        try: dialect=csv.Sniffer().sniff(sample,delimiters=",;\t")
        except: dialect=csv.excel
        r=csv.DictReader(io.StringIO(text),dialect=dialect)
        headers=[normalize_header(x) for x in (r.fieldnames or [])]
        rows=[]
        for row in r:
            rows.append({normalize_header(k):v for k,v in row.items() if k is not None})
        return headers, rows
    raise ValueError("Formato não suportado. Use .xlsx ou .csv")

def import_rows(rows, update_existing=False):
    # Carrega as chaves existentes uma única vez para evitar uma consulta
    # SELECT por linha da planilha. Isso deixa importações grandes muito mais rápidas.
    con=sqlite3.connect(DB); cur=con.cursor()
    dbcols=[r[1] for r in cur.execute("PRAGMA table_info(jogos)").fetchall()]
    allowed=set(dbcols)
    existing={}
    for rowid,data,hora,confronto,camp in cur.execute('SELECT rowid, CAST("Data" AS TEXT), CAST("Hora" AS TEXT), CAST("Confronto" AS TEXT), CAST("Campeonato" AS TEXT) FROM jogos'):
        key=(str(data or "").strip(),str(hora or "").strip(),str(confronto or "").strip(),str(camp or "").strip())
        existing[key]=rowid
    inserted=updated=dupes=invalid=0
    for raw in rows:
        row={k:raw.get(k) for k in allowed if k in raw}
        required=["Data","Confronto","Campeonato","Gols Casa FT","Gols Visitante FT"]
        if any(row.get(k) in (None,"") for k in required):
            invalid+=1; continue
        key=(str(row.get("Data","")).strip(),str(row.get("Hora","")).strip(),str(row.get("Confronto","")).strip(),str(row.get("Campeonato","")).strip())
        rowid=existing.get(key)
        if rowid is not None:
            if not update_existing:
                dupes+=1; continue
            sets=[c for c in row if c not in {"Data","Hora","Confronto","Campeonato"}]
            if sets:
                cur.execute('UPDATE jogos SET '+','.join('"'+c.replace('"','""')+'"=?' for c in sets)+' WHERE rowid=?',[row[c] for c in sets]+[rowid])
                updated+=1
            else:
                dupes+=1
        else:
            cols=list(row)
            cur.execute('INSERT INTO jogos ('+','.join('"'+c.replace('"','""')+'"' for c in cols)+') VALUES ('+','.join('?' for _ in cols)+')',[row[c] for c in cols])
            existing[key]=cur.lastrowid
            inserted+=1
    con.commit(); con.close()
    return {"inserted":inserted,"updated":updated,"duplicates":dupes,"invalid":invalid,"total_processed":len(rows)}

class Handler(BaseHTTPRequestHandler):
    def json(self,obj,code=200):
        b=json.dumps(json_safe(obj),ensure_ascii=False,default=str).encode(); self.send_response(code); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Access-Control-Allow-Origin","*"); self.send_header("Content-Length",str(len(b))); self.end_headers(); self.wfile.write(b)
    def do_POST(self):
        u=urlparse(self.path)
        try:
            n=int(self.headers.get("Content-Length","0") or 0); body=self.rfile.read(n)
            if u.path in ("/api/import/preview","/api/import"):
                filename,data=parse_upload(body,self.headers.get("Content-Type","")); headers,rows=read_rows(filename,data)
                if u.path=="/api/import/preview":
                    con=sqlite3.connect(DB); dbcols=[r[1] for r in con.execute("PRAGMA table_info(jogos)").fetchall()]; con.close(); matched=[h for h in headers if h in dbcols]
                    return self.json({"status":"ok","filename":filename,"rows":len(rows),"headers":headers,"matched":matched,"preview":json_safe(rows[:8])})
                q=parse_qs(u.query); update=q.get("update",["0"])[0] in ("1","true","True"); result=import_rows(rows,update); return self.json({"status":"ok","filename":filename,**result})
            return self.json({"status":"erro","msg":"Endpoint inválido"},404)
        except Exception as e: return self.json({"status":"erro","msg":str(e)},400)
    def do_GET(self):
        u=urlparse(self.path); p=parse_qs(u.query)
        try:
            if u.path=="/api/meta":
                con=sqlite3.connect(DB); champs=[r[0] for r in con.execute('SELECT DISTINCT "Campeonato" FROM jogos WHERE "Campeonato" IS NOT NULL AND TRIM("Campeonato")<>"" ORDER BY "Campeonato"').fetchall()]; total=con.execute("SELECT COUNT(*) FROM jogos").fetchone()[0]; con.close(); return self.json({"metrics":list(METRICS.keys()),"markets":list(MARKETS.keys()),"championships":champs,"total":total})
            if u.path=="/api/evaluate":
                strategies=json.loads(p.get("strategies",["[]"])[0]); market=p.get("market",["Over 1,5 FT"])[0]; stake=fnum(p.get("stake",["1"])[0]) or 1
                return self.json({"status":"ok","results":[evaluate(s,market,stake) for s in strategies]})
            if u.path in ("/", "/index.html"):
                b=(BASE/"index.html").read_bytes()
                self.send_response(200)
                self.send_header("Content-Type","text/html; charset=utf-8")
                self.send_header("Content-Length",str(len(b)))
                self.end_headers()
                self.wfile.write(b)
                return
            return self.json({"status":"erro","msg":"Não encontrado"},404)
        except Exception as e: return self.json({"status":"erro","msg":str(e)},500)
    def log_message(self,*args): pass

if __name__=="__main__":
    print(f"Banco: {DB}\nServidor Futebol Analyzer: porta {PORT}")
    HTTPServer(("0.0.0.0",PORT),Handler).serve_forever()
